import openpyxl
import re
from datetime import datetime, date, time
from openpyxl.utils.datetime import from_excel
from app.database import SessionLocal, ScheduleEntry, DutyAssignment

STOP_WORDS = ["часы", "смены", "больные", "отпуска", "выходные", "учёба", "норматив", "итого", "регистрация"]

HOMOGLYPHS = str.maketrans({
    'a': 'а', 'c': 'с', 'e': 'е', 'o': 'о', 'p': 'р', 'x': 'х', 'y': 'у',
    'A': 'А', 'B': 'В', 'C': 'С', 'E': 'Е', 'H': 'Н', 'K': 'К', 'M': 'М',
    'O': 'О', 'P': 'Р', 'T': 'Т', 'X': 'Х'
})

def get_surname_stem(full_or_surname: str) -> str:
    """Извлекает основу фамилии для точного сопоставления (Владимирова -> владимир, Блеч -> блеч)"""
    surname = full_or_surname.strip().split()[0].lower().translate(HOMOGLYPHS).replace('ё', 'е')
    for ending in ('ова', 'ева', 'ина', 'ая', 'яя', 'ов', 'ев', 'ин', 'ий', 'ый'):
        if surname.endswith(ending) and len(surname) - len(ending) >= 3:
            return surname[:-len(ending)]
    if surname.endswith(('а', 'я')) and len(surname) > 4:
        return surname[:-1]
    if surname.endswith('ко') and len(surname) > 4:
        return surname[:-1]
    return surname

def extract_date(val):
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if isinstance(val, (int, float)) and val > 40000:
        try:
            return from_excel(val).date()
        except Exception:
            pass
    if isinstance(val, str):
        s = val.strip()
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        m = re.search(r'(\d{1,2})[./](\d{1,2})[./](\d{4})', s)
        if m:
            try:
                return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                pass
    return None

def extract_shift(val):
    if val is None:
        return None, False, False

    if isinstance(val, (datetime, time)):
        h = val.hour
        if h in (7, 8, 9, 10, 12, 15):
            return f"{h:02d}:00", False, False
        return None, False, False

    if isinstance(val, float) and 0.0 < val < 1.0:
        h = int(round(val * 24))
        if h in (7, 8, 9, 10, 12, 15):
            return f"{h:02d}:00", False, False

    s = str(val).strip().upper().translate(HOMOGLYPHS)

    if s in ("О", "ОТ", "ОТПУСК"):
        return None, True, False
    if s in ("Б", "БОЛЬНИЧН", "БОЛЬНИЧНЫЙ"):
        return None, False, True
    if s in ("В", "У", "0", "-", "ВЫХОДНОЙ", "УВОЛЕН"):
        return None, False, False

    m = re.search(r'\b0?([789]|1[025]):00\b', s)
    if m:
        h = int(m.group(1))
        return f"{h:02d}:00", False, False

    m2 = re.search(r'[СC]\s*([789]|1[025])\b', s)
    if m2:
        h = int(m2.group(1))
        return f"{h:02d}:00", False, False

    return None, False, False

def parse_schedule_excel(file_path: str):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)

        sheet = None
        for name in wb.sheetnames:
            ws = wb[name]
            for row in ws.iter_rows(max_row=25, max_col=10, values_only=True):
                txt = " ".join([str(c) for c in row if c is not None]).lower()
                if any(k in txt for k in ["стрельников", "малахов", "пере(+)", "регистрация", "смены"]):
                    sheet = ws
                    break
            if sheet:
                break
        if not sheet:
            sheet = wb.active

        date_row_idx = None
        date_columns = {}

        for r_idx, row in enumerate(sheet.iter_rows(max_row=20, values_only=True), start=1):
            temp_dates = {}
            for c_idx, cell in enumerate(row):
                d = extract_date(cell)
                if d:
                    temp_dates[c_idx] = d
            if len(temp_dates) >= 5:
                date_row_idx = r_idx
                date_columns = temp_dates
                break

        if not date_row_idx:
            return False, f"Не удалось найти строку с датами на листе '{sheet.title}'."

        unique_shifts = {}
        known_employees = []
        current_row = date_row_idx + 1

        # 1. ЧТЕНИЕ ПЕРВОЙ ТАБЛИЦЫ (СОТРУДНИКИ И СМЕНЫ)
        while current_row <= sheet.max_row:
            row_cells = [cell.value for cell in sheet[current_row]]
            if not row_cells:
                current_row += 1
                continue

            emp_name = None
            is_stop = False
            for c in range(min(4, len(row_cells))):
                val = str(row_cells[c]).strip() if row_cells[c] is not None else ""
                val_lower = val.lower()
                if any(sw in val_lower for sw in STOP_WORDS):
                    is_stop = True
                    break
                if len(val) >= 5 and not val.replace('.', '').isdigit():
                    words = val.split()
                    if len(words) >= 2:
                        emp_name = f"{words[0]} {words[1]}"
                        break

            if is_stop:
                break

            if emp_name:
                if emp_name not in known_employees:
                    known_employees.append(emp_name)

                for col_idx, d_obj in date_columns.items():
                    if col_idx < len(row_cells):
                        shift, is_vac, is_sick = extract_shift(row_cells[col_idx])
                        if shift or is_vac or is_sick:
                            key = (emp_name, d_obj)
                            if key not in unique_shifts:
                                unique_shifts[key] = {
                                    "date": d_obj,
                                    "employee_name": emp_name,
                                    "shift_type": shift,
                                    "is_vacation": is_vac,
                                    "is_sick": is_sick
                                }
            current_row += 1

        # 2. ЧТЕНИЕ ВТОРОЙ ТАБЛИЦЫ (РЕГИСТРАЦИИ ПО ФАМИЛИЯМ)
        current_system = None
        unique_duties = {}

        while current_row <= sheet.max_row:
            row_cells = [cell.value for cell in sheet[current_row]]
            row_text = " ".join([str(c).strip() for c in row_cells[:4] if c is not None]).upper()

            if "НРД" in row_text: current_system = "НРД"
            elif "НКЦ" in row_text: current_system = "НКЦ"
            elif "МБ" in row_text: current_system = "МБ"

            slot = None
            if "УТРЕНН" in row_text or "07:00" in row_text: slot = "morning"
            elif "ДНЕВН" in row_text or "09:00" in row_text: slot = "day"
            elif "ВЕЧЕРН" in row_text or "18:00" in row_text: slot = "evening"

            if current_system and slot:
                for col_idx, d_obj in date_columns.items():
                    if col_idx < len(row_cells):
                        cell_raw = str(row_cells[col_idx] or "").strip()
                        if not cell_raw or cell_raw in ("0", "-") or cell_raw.lower() in ("в", "выходной"):
                            continue

                        clean_cell = cell_raw.translate(HOMOGLYPHS).replace('ё', 'е').lower()
                        # Извлекаем все слова из ячейки
                        words = re.findall(r'[а-яa-z]+', clean_cell)
                        matched = []

                        for w in words:
                            # Пропускаем одно- и двухбуквенные инициалы (А., В., И. и т.д.)
                            if len(w) < 3:
                                continue

                            w_stem = get_surname_stem(w)
                            for emp in known_employees:
                                emp_stem = get_surname_stem(emp)
                                emp_surname_exact = emp.split()[0].lower().translate(HOMOGLYPHS).replace('ё', 'е')
                                
                                # Сравнение строго по фамилии
                                if (w_stem == emp_stem or w == emp_surname_exact) and emp not in matched:
                                    matched.append(emp)
                                    break

                        # Резервный вариант, если сотрудника почему-то не было в списке смен
                        if not matched:
                            fallback = [w.capitalize() for w in words if len(w) >= 3]
                            if fallback:
                                matched = fallback

                        if matched:
                            key = (d_obj, current_system, slot)
                            unique_duties[key] = DutyAssignment(
                                date=d_obj,
                                system=current_system,
                                slot=slot,
                                employee_name=" / ".join(matched)
                            )

            current_row += 1

        db = SessionLocal()
        db.query(ScheduleEntry).delete()
        db.query(DutyAssignment).delete()

        for item in unique_shifts.values():
            db.add(ScheduleEntry(**item))

        for item in unique_duties.values():
            db.add(item)

        db.commit()
        db.close()
        return True, f"Успешно импортировано {len(unique_shifts)} смен и {len(unique_duties)} дежурств."
    except Exception as e:
        return False, f"Ошибка при обработке файла: {str(e)}"
